import glob, sys
from collections import Counter
import openpyxl

SILENCE = {'not addressed', 'none', 'none identified', 'n/a', 'na', 'not available', 'no'}
CRYPTO = {'MAC', 'Signature', 'TESLA', 'Certificate', 'Encryption', 'Blockchain'}
THIRD_PARTY = ('Research', 'Crowdsourced', 'Satellite', 'Mixed')


def find_xlsx(folder='data'):
    files = sorted(f for f in glob.glob(f'{folder}/*.xlsx') if not f.split('/')[-1].startswith('~$'))
    if not files:
        sys.exit(f'No .xlsx found in {folder}/ - put the coding workbook there and run again.')
    return files[0]


def norm(v, merge):
    if v is None or str(v).strip() == '':
        return None
    v = str(v).strip()
    return 'NONE' if merge and v.lower() in SILENCE else v


def load(path, merge=True):
    wb = openpyxl.load_workbook(path, data_only=True)
    coders = wb.sheetnames
    heads = [str(wb[coders[0]].cell(2, c).value) for c in range(4, 24)]
    codes = {s: {wb[s].cell(r, 3).value: [norm(wb[s].cell(r, c).value, merge) for c in range(4, 24)]
                 for r in range(3, wb[s].max_row + 1) if wb[s].cell(r, 3).value and wb[s].cell(r, 4).value}
             for s in coders}
    return coders, heads, codes


def key(head):
    return head.split(' (')[0].split(' /')[0]


def consensus(coders, codes):
    papers = sorted(set.intersection(*map(set, codes.values())), key=lambda p: int(p[1:]))
    cons, ties = {}, []
    for p in papers:
        row = []
        for i in range(20):
            votes = Counter(v for v in (codes[s][p][i] for s in coders) if v is not None).most_common()
            row.append(votes[0][0] if votes and (votes[0][1] >= 2 or len(votes) == 1) else 'TIE')
            if row[-1] == 'TIE':
                ties.append((p, i, [codes[s][p][i] for s in coders]))
        cons[p] = row
    return papers, cons, ties


def families(cons, heads, papers):
    fi = [key(h) for h in heads].index('Family')
    crypto = [p for p in papers if {x.strip() for x in cons[p][fi].split(',')} & CRYPTO]
    return crypto, [p for p in papers if p not in crypto]
