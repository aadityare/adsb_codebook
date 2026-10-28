import openpyxl


def anonymize(src, dst):
    wb = openpyxl.load_workbook(src, data_only=True)
    out = openpyxl.Workbook()
    out.remove(out.active)
    for i, ws in enumerate(wb.worksheets):
        new = out.create_sheet(f'Coder {chr(65 + i)}')
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=24, values_only=True):
            new.append(row)
    out.properties.creator = out.properties.lastModifiedBy = 'anonymous'
    out.save(dst)
    return dst
